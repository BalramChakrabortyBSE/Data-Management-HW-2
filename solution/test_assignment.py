"""Assignment-focused checks with an Excel scenario report."""

from __future__ import annotations

import re
import subprocess
import sys
import unittest
from pathlib import Path

import duckdb

ROOT = Path(__file__).resolve().parent.parent
WAREHOUSE = ROOT / "warehouse.duckdb"
SOURCE = ROOT / "data" / "f1.duckdb"
SOLUTION = ROOT / "solution"
REPORT = SOLUTION / "test_scenarios.xlsx"

SCENARIOS = {
    "test_00_clean_build": (
        "Parts 1–2: schema and loader",
        "Rebuild the warehouse from scratch using the submitted schema and load.",
        "The build command exits successfully and reports 5,105 fact_result rows.",
    ),
    "test_01_fact_grain": (
        "Part 1: fact table grain",
        "Verify one fact row per driver per race with the assignment's row count.",
        "fact_result contains 5,105 rows and no duplicate (race_key, driver_key) pairs.",
    ),
    "test_02_dimension_references": (
        "Part 1: star schema",
        "Verify every fact row references an existing race, driver, and team.",
        "There are no orphaned dimension keys.",
    ),
    "test_03_pit_stop_aggregation": (
        "Part 2: loader",
        "Check that stops are counted per entry without multiplying fact rows.",
        "The fact-table pit stop total matches the source and per-entry stop counts.",
    ),
    "test_04_assignment_check_values": (
        "Part 2: required validation values",
        "Check the required McLaren 2024 and Verstappen 2021 point totals.",
        "McLaren has 609 points in 2024; Max Verstappen has 388.5 in 2021.",
    ),
    "test_05_question_1": (
        "Part 3, Question 1",
        "Check the top-ten driver output, including wins, podiums, and races.",
        "The query returns the ten expected drivers with their required measures.",
    ),
    "test_06_question_2": (
        "Part 3, Question 2",
        "Check qualifying filters, minimum circuit history, averages, and ordering.",
        "Only classified grid starters are averaged; circuits have at least five races and results sort by the unrounded average.",
    ),
    "test_07_question_3": (
        "Part 3, Question 3",
        "Check that the star query returns non-duplicated 2024 team points and stops.",
        "McLaren is first with 609 points and 83 stops; Red Bull has 537 points and 85 stops.",
    ),
    "test_08_dashboard": (
        "Part 4: dashboard",
        "Check for two star-based charts, a season slider, and a findings headline.",
        "The dashboard has the required headline, season input, and two star-query charts.",
    ),
    "test_09_part_5_team_renames": (
        "Part 5, Question 1",
        "Check that the answer explains team-name totals and a way to group renamed teams.",
        "The answer explains the separate team rows and a shared team identifier/alias.",
    ),
    "test_10_part_5_fact_grain": (
        "Part 5, Question 2",
        "Check that the answer explains why pit stops should not be added as fact rows.",
        "The answer explains the grain mismatch and resulting duplication/double counting.",
    ),
    "test_11_submission_answers": (
        "Parts 3–4: submitted evidence",
        "Check that answers include the required explanation and dashboard screenshot.",
        "The answers contain the Monaco and join-grain explanations plus a dashboard image.",
    ),
}


class AssignmentTests(unittest.TestCase):
    def connection(self) -> duckdb.DuckDBPyConnection:
        return duckdb.connect(str(WAREHOUSE), read_only=True)

    def test_00_clean_build(self) -> None:
        """Rebuild from scratch and check the required fact row count."""
        result = subprocess.run(
            [sys.executable, "scripts/build.py"],
            cwd=ROOT,
            check=False,
            capture_output=True,
            text=True,
        )
        self.assertEqual(result.returncode, 0, result.stderr or result.stdout)
        self.assertIn("fact_result                5,105 rows", result.stdout)

    def test_01_fact_grain(self) -> None:
        """Check the fact table row count and composite grain."""
        with self.connection() as con:
            count, unique_grain = con.execute(
                """
                SELECT count(*), count(DISTINCT (race_key, driver_key))
                FROM fact_result
                """
            ).fetchone()
        self.assertEqual(count, 5105)
        self.assertEqual(unique_grain, count)

    def test_02_dimension_references(self) -> None:
        """Check all fact rows have matching dimension records."""
        with self.connection() as con:
            orphan_count = con.execute(
                """
                SELECT count(*)
                FROM fact_result f
                LEFT JOIN dim_race r USING (race_key)
                LEFT JOIN dim_driver d USING (driver_key)
                LEFT JOIN dim_team t USING (team_key)
                WHERE r.race_key IS NULL OR d.driver_key IS NULL OR t.team_key IS NULL
                """
            ).fetchone()[0]
        self.assertEqual(orphan_count, 0)

    def test_03_pit_stop_aggregation(self) -> None:
        """Compare per-entry warehouse pit stops with source stops."""
        with self.connection() as con:
            con.execute(f"ATTACH '{SOURCE}' AS src (READ_ONLY)")
            mismatch_count = con.execute(
                """
                WITH source_stops AS (
                    SELECT entry_id, count(*) AS stop_count
                    FROM src.pit_stops
                    GROUP BY entry_id
                )
                SELECT count(*)
                FROM fact_result f
                JOIN src.entries e
                  ON e.race_id = f.race_key AND e.driver_id = f.driver_key
                LEFT JOIN source_stops s USING (entry_id)
                WHERE f.pit_stops != coalesce(s.stop_count, 0)
                """
            ).fetchone()[0]
            warehouse_total, source_total = con.execute(
                """
                SELECT
                    (SELECT sum(pit_stops) FROM fact_result),
                    (SELECT count(*) FROM src.pit_stops)
                """
            ).fetchone()
        self.assertEqual(mismatch_count, 0)
        self.assertEqual(warehouse_total, source_total)

    def test_04_assignment_check_values(self) -> None:
        """Check the assignment's required point totals."""
        with self.connection() as con:
            mclaren_points = con.execute(
                """
                SELECT sum(f.points)
                FROM fact_result f
                JOIN dim_team t USING (team_key)
                JOIN dim_race r USING (race_key)
                WHERE t.team_name = 'McLaren' AND r.season = 2024
                """
            ).fetchone()[0]
            verstappen_points = con.execute(
                """
                SELECT sum(f.points)
                FROM fact_result f
                JOIN dim_driver d USING (driver_key)
                JOIN dim_race r USING (race_key)
                WHERE d.driver_name = 'Max Verstappen' AND r.season = 2021
                """
            ).fetchone()[0]
        self.assertEqual(mclaren_points, 609)
        self.assertEqual(verstappen_points, 388.5)

    def query_statements(self) -> list[str]:
        sql = (SOLUTION / "queries.sql").read_text()
        return [
            statement.strip()
            for statement in re.split(r";\s*", sql)
            if statement.strip() and re.search(r"\bSELECT\b", statement, re.I)
        ]

    def test_05_question_1(self) -> None:
        """Check the required top-ten driver query and measures."""
        statements = self.query_statements()
        self.assertEqual(len(statements), 3)
        with self.connection() as con:
            rows = con.execute(statements[0]).fetchall()
        self.assertEqual(
            rows,
            [
                ("Lewis Hamilton", 83, 148, 251),
                ("Max Verstappen", 71, 127, 233),
                ("Nico Rosberg", 20, 46, 59),
                ("Sebastian Vettel", 14, 60, 180),
                ("Lando Norris", 11, 44, 152),
                ("Valtteri Bottas", 10, 67, 228),
                ("Oscar Piastri", 9, 26, 70),
                ("Charles Leclerc", 8, 50, 173),
                ("Daniel Ricciardo", 8, 32, 207),
                ("Sergio Pérez", 6, 36, 226),
            ],
        )

    def test_06_question_2(self) -> None:
        """Check eligible circuit averages and descending full-precision order."""
        statements = self.query_statements()
        self.assertEqual(len(statements), 3)
        with self.connection() as con:
            rows = con.execute(statements[1]).fetchall()
        self.assertEqual(len(rows), 22)
        self.assertTrue(all(row[2] >= 5 for row in rows))
        averages = [row[3] for row in rows]
        self.assertEqual(averages, sorted(averages, reverse=True))
        monaco = next(row for row in rows if row[0] == "Circuit de Monaco")
        self.assertEqual(monaco[3], 1.125)
        normalized = re.sub(r"\s+", " ", statements[1].lower())
        self.assertIn("where f.is_classified and f.grid > 0", normalized)
        self.assertIn("order by avg_places_gained desc", normalized)

    def test_07_question_3(self) -> None:
        """Compare the 2024 star query with independently aggregated source data."""
        statements = self.query_statements()
        self.assertEqual(len(statements), 3)
        with self.connection() as con:
            con.execute(f"ATTACH '{SOURCE}' AS src (READ_ONLY)")
            rows = con.execute(statements[2]).fetchall()
            source_rows = con.execute(
                """
                WITH stops AS (
                    SELECT entry_id, count(*) AS pit_stops
                    FROM src.pit_stops
                    GROUP BY entry_id
                )
                SELECT t.name, sum(res.points), sum(coalesce(s.pit_stops, 0))
                FROM src.results res
                JOIN src.entries e USING (entry_id)
                JOIN src.races r USING (race_id)
                JOIN src.teams t USING (team_id)
                LEFT JOIN stops s USING (entry_id)
                WHERE r.season = 2024
                GROUP BY t.name
                ORDER BY sum(res.points) DESC
                """
            ).fetchall()
        self.assertEqual(rows, source_rows)
        self.assertEqual(rows[0], ("McLaren", 609.0, 83))
        self.assertIn(("Red Bull", 537.0, 85), rows)

    def test_08_dashboard(self) -> None:
        """Check dashboard slider, star-based charts, and findings headline."""
        dashboard = (SOLUTION / "dashboard.py").read_text()
        self.assertIn('st.sidebar.slider("Seasons"', dashboard)
        self.assertIn("McLaren scored the most points in 2024 with 609", dashboard)
        self.assertIn("st.bar_chart(", dashboard)
        self.assertIn("st.line_chart(", dashboard)
        self.assertGreaterEqual(dashboard.count("FROM fact_result"), 2)
        self.assertNotIn("src.", dashboard)

    def answer_part_5(self, question_number: int) -> str:
        answers = (SOLUTION / "answers.md").read_text()
        section = answers.split("## Part 5", maxsplit=1)[1]
        next_section = re.search(r"(?m)^## ", section)
        if next_section:
            section = section[: next_section.start()]
        question_pattern = rf"(?ms)^\s*{question_number}\.\s*(.*?)(?=^\s*[1-5]\.\s*|\Z)"
        match = re.search(question_pattern, section)
        if not match:
            return ""
        answer = match.group(1).strip()
        prompt = {
            1: "Toro Rosso, AlphaTauri and RB:",
            2: "One row per pit stop:",
        }[question_number]
        if answer.startswith(prompt):
            answer = answer[len(prompt) :].strip()
        return answer

    def test_09_part_5_team_renames(self) -> None:
        """Check the team rename and grouping explanation."""
        answer = self.answer_part_5(1)
        self.assertTrue(
            answer,
            "Part 5 Question 1 contains the prompt but no answer.",
        )
        self.assertRegex(answer.lower(), r"team|points")
        self.assertRegex(answer.lower(), r"team.?id|alias|group|rename")

    def test_10_part_5_fact_grain(self) -> None:
        """Check the pit-stop fact-grain explanation."""
        answer = self.answer_part_5(2)
        self.assertTrue(
            answer,
            "Part 5 Question 2 contains the prompt but no answer.",
        )
        self.assertRegex(answer.lower(), r"grain|pit stop")
        self.assertRegex(answer.lower(), r"duplicate|double count|repeat|inflate")

    def test_11_submission_answers(self) -> None:
        """Check Monaco/join-grain explanations and screenshot evidence."""
        answers = (SOLUTION / "answers.md").read_text()
        self.assertIn("Monaco averages 1.125", answers)
        self.assertIn("grain", answers.lower())
        self.assertRegex(answers, r"!\[[^\]]*\]\(dashboard\.png\)")
        self.assertTrue((SOLUTION / "dashboard.png").is_file())


class ScenarioResult(unittest.TextTestResult):
    def __init__(self, *args: object, **kwargs: object) -> None:
        super().__init__(*args, **kwargs)
        self.scenario_results: dict[str, tuple[str, str]] = {}

    def addSuccess(self, test: unittest.case.TestCase) -> None:
        super().addSuccess(test)
        self.scenario_results[test._testMethodName] = (
            "Pass",
            "Automated checks passed.",
        )

    def addFailure(
        self,
        test: unittest.case.TestCase,
        err: tuple[type[BaseException], BaseException, object],
    ) -> None:
        super().addFailure(test, err)
        self.scenario_results[test._testMethodName] = (
            "Failed",
            self._failure_detail(err),
        )

    def addError(
        self,
        test: unittest.case.TestCase,
        err: tuple[type[BaseException], BaseException, object],
    ) -> None:
        super().addError(test, err)
        self.scenario_results[test._testMethodName] = (
            "Failed",
            self._failure_detail(err),
        )

    @staticmethod
    def _failure_detail(err: tuple[type[BaseException], BaseException, object]) -> str:
        return str(err[1])[:1000]


class ScenarioRunner(unittest.TextTestRunner):
    resultclass = ScenarioResult


def write_workbook(results: ScenarioResult) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Test Scenarios"
    sheet.append(
        [
            "ID",
            "Assignment requirement",
            "Scenario tested",
            "Expected result",
            "Actual result",
            "Status",
        ]
    )
    for index, (test_name, (area, scenario, expected)) in enumerate(
        SCENARIOS.items(), start=1
    ):
        status, actual = results.scenario_results.get(
            test_name, ("Failed", "Test did not complete.")
        )
        sheet.append([f"TC-{index:02d}", area, scenario, expected, actual, status])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    pass_fill = PatternFill("solid", fgColor="E2F0D9")
    failed_fill = PatternFill("solid", fgColor="FCE4D6")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        row[5].fill = pass_fill if row[5].value == "Pass" else failed_fill
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {
        "A": 10,
        "B": 27,
        "C": 48,
        "D": 60,
        "E": 60,
        "F": 12,
    }.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 30
    workbook.save(REPORT)


def main() -> int:
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(AssignmentTests)
    result = ScenarioRunner(verbosity=2).run(suite)
    write_workbook(result)
    print(f"\nExcel scenario report: {REPORT}")
    return 0 if result.wasSuccessful() else 1


if __name__ == "__main__":
    raise SystemExit(main())
